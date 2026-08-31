"""Fill only the visible NetPay certification questions requested."""
import glob
import re
import shutil

import openpyxl
from openpyxl.styles import Alignment

SRC_MATCH = r"c:\Users\rogel\Downloads\Resultados_Certificaci*San_Martin*v4.4.1*(1).xlsx"
DST = r"c:\Users\rogel\Downloads\Resultados_Certificacion_NetPay_RESPONDIDO.xlsx"
DST_COPY = r"c:\desarrollo\sanmartin\netpay-front\Resultados_Certificacion_NetPay_RESPONDIDO.xlsx"

candidates = [f for f in glob.glob(SRC_MATCH) if "RESPONDIDO" not in f.upper()]
if not candidates:
    raise FileNotFoundError("No se encontró el Excel (1) en Downloads")
SRC = candidates[0]

# Solo las preguntas visibles pedidas.
ANSWERS = {
    1: (
        "El usuario entra a Odoo con su cuenta de la empresa. Ese login no genera el token de NetPay ni está ligado a un solo storeId. "
        "El token OAuth lo obtiene nuestro servidor (api-netpay) al hablar con NetPay. "
        "El panel administrativo tiene login propio, también separado de NetPay."
    ),
    2: (
        "Sí. En cada cobro se busca la terminal del usuario Odoo y la empresa: storeId y número de serie (folio de terminal). "
        "Un usuario puede tener distinta terminal por empresa. No hay un storeId fijo para todo el sistema."
    ),
    3: (
        "Se cambia de pruebas a producción con configuración del servidor: URL de NetPay, credenciales del token, URL del webhook y claves internas. "
        "Cada ambiente tiene sus propios valores; no se mezclan."
    ),
    4: (
        "El servidor pide el token a oauth-service/oauth/token. No usamos refresh token. "
        "Se genera al primer uso o si NetPay responde 401/403; entonces se pide uno nuevo y se reintenta la misma operación. "
        "No hay un job que lo renueve cada 6 u 11 horas; el vencimiento se cubre con ese reintento."
    ),
    5: (
        "No. En certificación y producción el webhook es público con HTTPS y certificado SSL (compatible con la terminal Android/Pax). "
        "La ruta es POST /transactions/terminal-response en nuestro servidor. En desarrollo interno puede usarse localhost."
    ),
    6: (
        "Usamos un webhook central en nuestro servidor. NetPay avisa ahí; se guarda la transacción y se libera la espera hacia Odoo. "
        "No reenviamos a un servidor local de sucursal. Si falla el proceso, se registra en bitácora y el cobro termina en error o timeout."
    ),
    9: (
        "Sí. El backend espera la respuesta de la terminal hasta 180 segundos (3 minutos) y consulta la base cada 500 ms por si ya llegó el webhook. "
        "Si no hay respuesta, se marca timeout. No hay cashback (no aplica espera de 5 minutos)."
    ),
    10: (
        "Odoo muestra una espera y bloquea la pantalla del cobro hasta que llegue la respuesta o el timeout del servidor. "
        "No hay cronómetro de 3 minutos en pantalla. Si el cliente termina antes, la respuesta llega antes. No usamos cashback."
    ),
    11: (
        "Sí. Antes de enviar el cobro a NetPay ya existe un registro con folioNumber único. "
        "Si la terminal no responde, el registro queda pendiente y se puede consultar o reimprimir por folio en el panel. "
        "Odoo no registra el pago hasta que el cobro se complete."
    ),
    12: (
        "No. Solo usamos las APIs documentadas de Smart Accounts: venta, cancelación y reimpresión. "
        "Nuestros endpoints propios solo orquestan la espera y el registro en Odoo. No hay APIs nuevas de NetPay sin contexto de integraciones."
    ),
    13: (
        "El folio se genera antes de enviar la venta y se manda en folioNumber. "
        "No se crea después de procesar el cobro en la terminal."
    ),
    14: (
        "Formato: id de la orden de pago + id interno de la transacción + timestamp. "
        "Ejemplo: 12345-678-1723488123456. No es un consecutivo numérico simple."
    ),
    15: (
        "Sí. El folio es único por cada intento de cobro en todo el sistema. "
        "No se repite entre sucursales, usuarios ni reintentos."
    ),
    16: (
        "No se reutiliza. Una declinación conserva su folio. "
        "Un nuevo intento genera otro folioNumber."
    ),
    17: (
        "Validamos responseCode: 00 es aprobado. "
        "En venta suele venir transType A; en cancelación en terminal, transType V. "
        "Cualquier código distinto de 00 es error, declinación o cancelación."
    ),
    18: (
        "Sí. Las declinaciones se guardan (código, orderId si viene, importe, folio, tarjeta, mensaje y si se aplicó o no). "
        "En el panel aparecen como declinadas."
    ),
    19: (
        "Sí. Se mantiene el voucher oficial de NetPay en la terminal Pax. "
        "No deshabilitamos la impresión oficial ni usamos ticket personalizado."
    ),
    23: "No. Solo se imprime en la terminal NetPay (venta, reimpresión y cancelación). No hay impresora propia adicional.",
    24: (
        "Sí. En cada venta enviamos traceability (venta, orden de pago, usuario, empresa, storeId, folio, monto). "
        "NetPay lo regresa en el callback para correlacionar. Esos datos no se almacenan en NetPay."
    ),
    25: (
        "Sí. Hay bitácora en servidor (peticiones, respuesta de terminal y registro en Odoo) y el panel muestra movimientos con detalle. "
        "No es solo log local en la PC del cajero."
    ),
    26: (
        "Computadora Windows/navegador: Odoo y el panel admin. "
        "La tarjeta se cobra en terminal Pax Smart PinPad. No es app Android propia ni kiosco."
    ),
    27: "No. La terminal usa la app NetPay estándar e internet. Sin VPN ni otra aplicación adicional.",
    28: (
        "Sí, hay firewall corporativo en sucursal. "
        "La terminal debe poder salir por HTTPS a NetPay y el webhook debe ser alcanzable desde internet por HTTPS."
    ),
    29: (
        "Si se pierde una transacción, se puede reimprimir desde el panel por folioNumber (folioId) o por orderId, solo como consulta. "
        "No automatizamos la reimpresión para crear el pago en Odoo cuando la venta no existe."
    ),
    30: (
        "No. Ante una declinación no se lanza reimpresión automática. "
        "La reimpresión automática es solo después de un pago aplicado en Odoo o de una cancelación exitosa."
    ),
    31: (
        "No. El estatus final se toma del webhook y de lo guardado en base (responseCode / tipo_respuesta). "
        "La reimpresión no se usa para definir reversos."
    ),
    32: (
        "Sí. Odoo bloquea la pantalla mientras espera el cobro. "
        "En el panel se deshabilita la acción de la fila en curso. "
        "Además el servidor bloquea por número de serie de la terminal: si otra PC intenta cobro, cancelación o reimpresión a la misma Pax, recibe que está ocupada hasta que termine o pasen ~3 minutos."
    ),
    33: "No. Es operación asistida por cajero. No es kiosco ni autoatendido.",
    49: (
        "Sí. Si el webhook recibe y procesa bien, responde HTTP 200. "
        "Si la petición es inválida o hay error interno, responde 4xx o 5xx."
    ),
    57: (
        "Declinada: código distinto de 00, folioNumber, folio de venta, importe, tarjeta, storeId, usuario, empresa, "
        "mensaje de la terminal y orderId si viene. Queda como no aplicada."
    ),
    58: (
        "Aprobada: responseCode 00, orderId, folioNumber, importe, tarjeta, storeId, usuario y empresa. "
        "Cuando Odoo crea el pago se marca aplicada y se guarda el id de ese pago."
    ),
    59: (
        "No hay catálogo aparte de reversos. "
        "Pendiente: sin código de respuesta. Si se canceló en terminal, queda el mensaje. "
        "No consultamos reverso por reimpresión."
    ),
    60: (
        "De la reimpresión guardamos tipo reimpresión, folioId u orderId usado y el resultado. "
        "Queda ligada al cobro original."
    ),
    61: (
        "De la cancelación guardamos el orderId de la venta, importe, storeId y el código de respuesta. "
        "La venta original se marca como devuelta. En el pago de Odoo se agrega una nota en el chatter."
    ),
    62: (
        "Sí. Lo que envía NetPay se persiste y se usa en movimientos, devoluciones y reimpresiones. "
        "Validar la conciliación del día es responsabilidad del comercio."
    ),
}

wb = openpyxl.load_workbook(SRC)
ws = wb["Cuestionario"]
answer_col = 6
first_row = 4
last_row = 65
wrap = Alignment(wrap_text=True, vertical="top")


def parse_qid(row: int, raw) -> int | None:
    if isinstance(raw, (int, float)):
        return int(raw)
    if raw is None:
        return row - first_row + 1 if first_row <= row <= last_row else None
    match = re.search(r",(\d+\.?\d*)\)$", str(raw))
    if match:
        return int(float(match.group(1)))
    return row - first_row + 1 if first_row <= row <= last_row else None


filled = []
for r in range(first_row, last_row + 1):
    qid = parse_qid(r, ws.cell(r, 2).value)
    if qid not in ANSWERS:
        continue
    cell = ws.cell(r, answer_col)
    cell.value = ANSWERS[qid]
    cell.alignment = wrap
    filled.append(qid)

ws.column_dimensions["F"].width = 80
wb.save(DST)
shutil.copy2(DST, DST_COPY)
print(f"Source: {SRC}")
print(f"Saved: {DST}")
print(f"Copy: {DST_COPY}")
print(f"Filled {len(filled)}: {filled}")
